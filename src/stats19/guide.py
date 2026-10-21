"""Parse the DfT road safety open data guide (.xlsx) into lookup rows.

Sheets expected (2025 guide):
- ``<year>_code_list``: table | field name | code/format | label | note
- ``<year1>_to_<year2>_conversion``: old variable/value/label -> new variable/value/label
"""

from __future__ import annotations

import csv
from collections.abc import Iterator
from dataclasses import dataclass
from importlib import resources
from pathlib import Path

from openpyxl import load_workbook

# Guide field names that differ from the CSV header (see FU-002).
FIELD_ALIASES: dict[str, str] = {
    "enhanced_collision_severity": "enhanced_severity_collision",
    "did_police_officer_attend_scene_of_collision": "did_police_officer_attend_scene_of_accident",
    "lsoa_of_collision_location": "lsoa_of_accident_location",
}


class GuideError(ValueError):
    """The guide workbook doesn't have the expected layout."""


@dataclass(frozen=True)
class LookupCode:
    table_name: str
    field_name: str
    code: str
    label: str
    note: str | None
    source: str = "dft_guide"


@dataclass(frozen=True)
class Conversion:
    table_name: str
    old_field: str
    old_code: str
    old_label: str
    new_field: str
    new_code: str
    new_label: str


@dataclass(frozen=True)
class Guide:
    codes: list[LookupCode]
    field_notes: list[tuple[str, str, str]]
    conversions: list[Conversion]


def _text(value: object) -> str:
    """Cell value as trimmed text; whole floats lose their '.0'."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _field(value: object) -> str:
    name = _text(value).lower()
    return FIELD_ALIASES.get(name, name)


def _rows(sheet) -> Iterator[list[str]]:
    rows = sheet.iter_rows(values_only=True)
    next(rows, None)  # header
    for row in rows:
        cells = [_text(v) for v in row]
        if any(cells):
            yield cells


def _sheet(workbook, suffix: str):
    for name in workbook.sheetnames:
        if name.endswith(suffix):
            return workbook[name]
    raise GuideError(f"no sheet ending in '{suffix}' (found {workbook.sheetnames})")


def _parse_codes(sheet) -> tuple[list[LookupCode], list[tuple[str, str, str]]]:
    codes: list[LookupCode] = []
    notes: list[tuple[str, str, str]] = []
    seen: set[tuple[str, str, str]] = set()
    for cells in _rows(sheet):
        table, field, code, label, note = (cells + [""] * 5)[:5]
        table, field = table.lower(), _field(field)
        if not code:
            if note:
                notes.append((table, field, note))
            continue
        key = (table, field, code)
        if key in seen:
            continue
        seen.add(key)
        codes.append(LookupCode(table, field, code, label, note or None))
    return codes, notes


def _parse_conversions(sheet) -> list[Conversion]:
    out: list[Conversion] = []
    for cells in _rows(sheet):
        table, old_f, old_c, old_l, new_f, new_c, new_l = (cells + [""] * 7)[:7]
        if not old_c or old_c == "-":
            continue
        out.append(
            Conversion(table.lower(), _field(old_f), old_c, old_l, _field(new_f), new_c, new_l)
        )
    return out


def parse_guide(path: Path) -> Guide:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        codes, notes = _parse_codes(_sheet(workbook, "code_list"))
        conversions = _parse_conversions(_sheet(workbook, "conversion"))
    finally:
        workbook.close()
    return Guide(codes=codes, field_notes=notes, conversions=conversions)


def load_supplement() -> list[LookupCode]:
    """Labels the DfT guide lacks (FU-009), bundled as sql/lookup_supplement.csv."""
    text = resources.files("stats19").joinpath("sql", "lookup_supplement.csv").read_text("utf-8")
    return [
        LookupCode(
            r["table_name"],
            r["field_name"],
            r["code"],
            r["label"],
            r["note"] or None,
            "supplement",
        )
        for r in csv.DictReader(text.splitlines())
    ]


def merge_supplement(guide: list[LookupCode], supplement: list[LookupCode]) -> list[LookupCode]:
    """Append supplement codes the guide doesn't define; the guide always wins."""
    known = {(c.table_name, c.field_name, c.code) for c in guide}
    extra = [c for c in supplement if (c.table_name, c.field_name, c.code) not in known]
    return [*guide, *extra]
