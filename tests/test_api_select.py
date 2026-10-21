"""Selection endpoints: summary, paged list, exports, filter description."""

import csv
import io
import json
import zipfile

import pytest
from openpyxl import load_workbook

pytestmark = pytest.mark.integration

# Fixture data (tests/test_pipeline_integration.py): collisions 1979A and 1980A survive,
# both severity code 1 (the default code in fixtures) -> counted as fatal.


def _f(**kw) -> dict:
    return {"f": json.dumps(kw)} if kw else {}


def test_summary_totals_and_breakdowns(api):
    d = api.get("/api/summary").json()["data"]

    assert d["totals"]["collisions"] == 2
    assert d["totals"]["fatal"] == 2 and d["totals"]["ksi"] == 2
    assert d["totals"]["ksi_adjusted"] == 2  # fatal collisions count 1 whatever the adjustment
    assert d["totals"]["casualties"] == 3
    assert [r["year"] for r in d["by_year"]] == [1979, 1980]
    assert d["by_year"][0] == {"year": 1979, "fatal": 1, "serious": 0, "slight": 0}
    light = d["breakdowns"]["light_conditions"]
    assert light == [{"code": 1, "label": "One", "n": 2}]
    assert sum(c["n"] for c in d["by_dow_hour"]) == 2
    assert {"group", "fatal", "serious", "slight"} <= set(d["casualties_by_group"][0])


def test_summary_respects_filter(api):
    d = api.get("/api/summary", params=_f(year_from=1980)).json()["data"]

    assert d["totals"]["collisions"] == 1
    assert [r["year"] for r in d["by_year"]] == [1980]


def test_summary_empty_selection(api):
    d = api.get("/api/summary", params=_f(severity=[3])).json()["data"]

    assert d["totals"]["collisions"] == 0
    assert d["by_year"] == []


def test_collision_list_pages_and_sorts(api):
    body = api.get("/api/collisions", params={"page_size": 1, "sort": "date_desc"}).json()

    assert body["meta"] == {"total": 2, "page": 1, "page_size": 1}
    row = body["data"][0]
    assert row["collision_index"] == "1980A"
    assert {
        "collision_date",
        "severity",
        "light_conditions_label",
        "number_of_casualties",
        "lon",
        "lat",
    } <= set(row)

    second = api.get(
        "/api/collisions", params={"page_size": 1, "page": 2, "sort": "date_desc"}
    ).json()
    assert second["data"][0]["collision_index"] == "1979A"


def test_collision_list_rejects_unknown_sort(api):
    assert api.get("/api/collisions", params={"sort": "collision_index; drop"}).status_code == 422


def test_csv_zip_export(api):
    r = api.get("/api/export/collisions.zip", params=_f(year_from=1979))

    assert r.status_code == 200
    assert r.headers["content-type"] == "application/zip"
    assert "attachment" in r.headers["content-disposition"]
    z = zipfile.ZipFile(io.BytesIO(r.content))
    assert sorted(z.namelist()) == [
        "README.txt",
        "casualties.csv",
        "collisions.csv",
        "vehicles.csv",
    ]
    rows = list(csv.DictReader(io.StringIO(z.read("collisions.csv").decode("utf-8-sig"))))
    assert sorted(r["collision_index"] for r in rows) == ["1979A", "1980A"]
    assert "light_conditions_label" in rows[0] and "geom_bng" not in rows[0]
    vehicles = list(csv.DictReader(io.StringIO(z.read("vehicles.csv").decode("utf-8-sig"))))
    assert len(vehicles) == 3
    readme = z.read("README.txt").decode()
    assert "Years: 1979–1980" in readme or "From 1979" in readme
    assert "Open Government Licence" in readme


def test_xlsx_export(api):
    r = api.get("/api/export/collisions.xlsx")

    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    assert wb.sheetnames == ["About", "Collisions", "Vehicles", "Casualties"]
    rows = list(wb["Collisions"].iter_rows(values_only=True))
    assert rows[0][0] == "collision_index"
    assert len(rows) == 3  # header + 2


def test_xlsx_dates_and_times_are_typed(api):
    import datetime as dt

    r = api.get("/api/export/collisions.xlsx")
    wb = load_workbook(io.BytesIO(r.content))  # full mode: number formats are read
    ws = wb["Collisions"]
    head = [c.value for c in ws[1]]
    row = {k: c for k, c in zip(head, ws[2], strict=True)}

    assert isinstance(row["collision_date"].value, dt.datetime)
    assert row["collision_date"].number_format == "yyyy-mm-dd"
    assert isinstance(row["collision_time"].value, dt.time)
    assert row["collision_time"].number_format == "hh:mm"


def test_export_too_large_is_refused(api, monkeypatch):
    import stats19.api.export as export

    monkeypatch.setattr(export, "MAX_EXPORT_COLLISIONS", 1)
    r = api.get("/api/export/collisions.zip")

    assert r.status_code == 413
    assert r.json()["error"]["code"] == "too_large"


def test_describe_filter(api):
    d = api.get(
        "/api/describe",
        params=_f(
            year_from=2021,
            year_to=2025,
            severity=[1, 2],
            involves=["motorcycle"],
            fields=[{"table": "collision", "field": "light_conditions", "values": [1, 2]}],
            area={"type": "radius", "lon": -3.53, "lat": 50.72, "radius_m": 500},
        ),
    ).json()["data"]

    assert "Years: 2021–2025" in d
    assert "Severity: Fatal, Serious" in d
    assert "Road users involved: Motorcycle" in d
    assert "Light conditions: One, Two" in d
    assert "Area: within 500 m of 50.72000, -3.53000" in d


def test_summary_too_large_is_refused(api, monkeypatch):
    import stats19.api.summary as summary

    monkeypatch.setattr(summary, "MAX_SUMMARY_COLLISIONS", 1)
    r = api.get("/api/summary")

    assert r.status_code == 413
    assert r.json()["error"]["code"] == "too_large"
