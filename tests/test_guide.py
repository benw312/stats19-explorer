import pytest
from openpyxl import Workbook

from stats19.guide import GuideError, parse_guide


def _write_guide(path, code_rows, conversion_rows=None):
    wb = Workbook()
    codes = wb.active
    codes.title = "2024_code_list"
    codes.append(["table", "field name", "code/format", "label", "note"])
    for row in code_rows:
        codes.append(row)
    conv = wb.create_sheet("2011_to_2024_conversion")
    conv.append(
        [
            "table",
            "2011 format variable",
            "2011 format value",
            "2011 format label",
            "2024 format variable",
            "2024 format value",
            "2024 format label",
        ]
    )
    for row in conversion_rows or []:
        conv.append(row)
    wb.save(path)
    return path


def test_parse_guide_returns_code_rows_with_text_codes(tmp_path):
    path = _write_guide(
        tmp_path / "g.xlsx",
        [
            ["collision", "collision_index", None, None, "unique value"],
            ["collision", "collision_severity", 1, "Fatal", None],
            ["collision", "collision_severity", 2.0, "Serious", None],
            ["Vehicle ", " Vehicle_Type ", "-1", "Data missing or out of range", None],
        ],
    )

    guide = parse_guide(path)

    assert [(c.table_name, c.field_name, c.code, c.label) for c in guide.codes] == [
        ("collision", "collision_severity", "1", "Fatal"),
        ("collision", "collision_severity", "2", "Serious"),
        ("vehicle", "vehicle_type", "-1", "Data missing or out of range"),
    ]
    assert guide.field_notes == [("collision", "collision_index", "unique value")]


def test_parse_guide_applies_field_aliases(tmp_path):
    path = _write_guide(
        tmp_path / "g.xlsx",
        [["collision", "enhanced_collision_severity", 5, "Very Serious", None]],
    )

    guide = parse_guide(path)

    assert guide.codes[0].field_name == "enhanced_severity_collision"


def test_parse_guide_skips_duplicate_codes_keeping_first(tmp_path):
    path = _write_guide(
        tmp_path / "g.xlsx",
        [
            ["casualty", "sex_of_casualty", 1, "Male", None],
            ["casualty", "sex_of_casualty", 1, "Male (dup)", None],
        ],
    )

    guide = parse_guide(path)

    assert [c.label for c in guide.codes] == ["Male"]


def test_parse_guide_reads_conversion_sheet(tmp_path):
    path = _write_guide(
        tmp_path / "g.xlsx",
        [],
        [
            [
                "collision",
                "junction_detail_historic",
                3,
                "T, Y or staggered junction",
                "junction_detail",
                13,
                "T or staggered junction",
            ],
            ["vehicle", "journey_purpose_of_driver_historic", "-", "-", "x", 8, "Emergency"],
        ],
    )

    guide = parse_guide(path)

    assert len(guide.conversions) == 1
    conv = guide.conversions[0]
    assert (conv.old_field, conv.old_code, conv.new_field, conv.new_code) == (
        "junction_detail_historic",
        "3",
        "junction_detail",
        "13",
    )


def test_parse_guide_rejects_unexpected_layout(tmp_path):
    wb = Workbook()
    wb.active.title = "something_else"
    wb.save(tmp_path / "bad.xlsx")

    with pytest.raises(GuideError, match="code_list"):
        parse_guide(tmp_path / "bad.xlsx")


# ---- supplementary labels (FU-009)

from stats19.guide import LookupCode, load_supplement, merge_supplement  # noqa: E402


def _code(table, field, code, label, source="dft_guide"):
    return LookupCode(table, field, code, label, None, source)


def test_merge_supplement_adds_missing_codes_with_provenance():
    guide = [_code("casualty", "casualty_type", "1", "Cyclist")]
    supplement = [_code("casualty", "casualty_type", "33", "PPT rider", "supplement")]

    merged = merge_supplement(guide, supplement)

    assert [(c.code, c.source) for c in merged] == [("1", "dft_guide"), ("33", "supplement")]


def test_merge_supplement_guide_wins_on_conflict():
    guide = [_code("casualty", "casualty_type", "33", "Official label")]
    supplement = [_code("casualty", "casualty_type", "33", "Our label", "supplement")]

    merged = merge_supplement(guide, supplement)

    assert [(c.label, c.source) for c in merged] == [("Official label", "dft_guide")]


def test_parsed_guide_codes_are_marked_dft_guide(tmp_path):
    path = _write_guide(
        tmp_path / "g.xlsx", [["collision", "collision_severity", 1, "Fatal", None]]
    )

    assert parse_guide(path).codes[0].source == "dft_guide"


def test_bundled_supplement_loads_and_cites_evidence():
    rows = load_supplement()

    keys = {(r.table_name, r.field_name, r.code) for r in rows}
    assert ("casualty", "casualty_type", "33") in keys
    assert ("collision", "collision_injury_based", "-1") in keys
    assert ("casualty", "casualty_injury_based", "-1") in keys
    assert all(r.source == "supplement" for r in rows)
    assert all(r.note and "FU-009" in r.note for r in rows), (
        "every supplement row cites its evidence"
    )


def test_lsoa_guide_name_aliased_to_csv_header(tmp_path):
    path = _write_guide(
        tmp_path / "g.xlsx", [["collision", "lsoa_of_collision_location", None, None, "LSOA code"]]
    )

    assert parse_guide(path).field_notes == [
        ("collision", "lsoa_of_accident_location", "LSOA code")
    ]
